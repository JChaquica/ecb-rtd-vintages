"""Reopen a sample of the workbooks written by get_data.py (code/workbooks.py) and
check them against the tables they were written from.

Why this check matters: the workbooks are what most users will open, and they
are a second copy of the data in another format. If a value were dropped,
shifted by a row or rounded on its way into Excel, no other check would show it.

For every series two workbooks are opened: the latest vintage and one earlier
vintage, drawn at random with a fixed seed so that a rerun opens the same ones.
Every date and every value is compared with the same vintage in
data/processed/raw_vintages. Two values count as equal when they agree to 14
significant digits. Excel stores 15, and a few of the ECB's values from 2024
onward carry 17, so those cannot come back exactly.

Output: output/logs/05_workbooks.txt

    python checks/05_workbooks.py
"""

from __future__ import annotations

import csv
import datetime
import os
import random
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(os.path.abspath(__file__)).parents[1] / "code"))

from common import (LOG_DIR, ROOT, TABLE_DIR, VARIABLE_LIST, WORKBOOK_DIR,  # noqa: E402
                    require_full_build)

LOG_PATH = LOG_DIR / "05_workbooks.txt"

SEED = 20261005
# Largest relative difference between two values that still counts as equal.
TOLERANCE = 1e-14


def observation_date(period: str) -> datetime.date:
    """First day of the period, as code/workbooks.py writes it."""
    period = period.strip()
    if "-Q" in period:
        year, quarter = period.split("-Q")
        return datetime.date(int(year), 3 * (int(quarter) - 1) + 1, 1)
    if "-" in period:
        year, month = period.split("-")
        return datetime.date(int(year), int(month), 1)
    return datetime.date(int(period), 1, 1)


def raw_vintage(key: str, date: str) -> list[tuple[datetime.date, float | None]]:
    """(observation_date, value) pairs of one vintage, from its first value to its last."""
    with (TABLE_DIR / f"{key}.csv").open(encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        column = next(reader).index(date)
        rows = [(observation_date(r[0]), r[column].strip()) for r in reader]
    filled = [i for i, (_, value) in enumerate(rows) if value]
    if not filled:
        return []
    return [(d, float(value) if value else None)
            for d, value in rows[filled[0]:filled[-1] + 1]]


def workbook_vintage(path: Path) -> list[tuple]:
    """The same pairs as a workbook holds them (its second sheet, below the header)."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book[book.sheetnames[-1]].iter_rows(values_only=True))[1:]
    book.close()
    return [(d.date() if isinstance(d, datetime.datetime) else d, value) for d, value in rows]


def main() -> None:
    # The tables must be complete; the workbooks are what is being checked, so a
    # series without them is reported below rather than stopping the check here.
    require_full_build()
    if not any(WORKBOOK_DIR.glob("*/*/vintage_*.xlsx")):
        sys.exit(f"no workbooks in {WORKBOOK_DIR.relative_to(ROOT).as_posix()}/. "
                 f"Run  python get_data.py --all  first.")
    with VARIABLE_LIST.open(encoding="utf-8-sig", newline="") as f:
        keys = {(row["Folder"], row["Variable"]): row["ECB series key"]
                for row in csv.DictReader(f)}

    rng = random.Random(SEED)
    books = dates = values = exact = single = 0
    worst = (0.0, "")
    bad = []

    for (folder, variable), key in sorted(keys.items()):
        files = sorted((WORKBOOK_DIR / folder / variable).glob("vintage_*.xlsx"))
        if not files:
            bad.append(f"{key}: no workbooks")
            continue
        chosen = [files[-1]]
        if len(files) > 1:
            chosen.append(rng.choice(files[:-1]))
        else:
            single += 1
        for path in sorted(chosen):
            date = path.stem.replace("vintage_", "")
            want, got = raw_vintage(key, date), workbook_vintage(path)
            books += 1
            if len(want) != len(got):
                bad.append(f"{key} {date}: {len(want)} rows rebuilt, {len(got)} in workbook")
                continue
            for (d1, v1), (d2, v2) in zip(want, got):
                dates += 1
                if d1 != d2:
                    bad.append(f"{key} {date}: date {d1} vs {d2}")
                    break
                if v1 is None and v2 is None:
                    continue
                values += 1
                if v1 is None or v2 is None:
                    bad.append(f"{key} {date} {d1}: {v1} vs {v2}")
                    break
                if v1 == v2:
                    exact += 1
                    continue
                relative = abs(v1 - v2) / max(abs(v1), abs(v2))
                if relative > worst[0]:
                    worst = (relative, f"{key} {date} {d1}: {v1!r} vs {v2!r}")
                if relative > TOLERANCE:
                    bad.append(f"{key} {date} {d1}: {v1!r} vs {v2!r} (relative {relative:.2e})")
                    break

    lines = [f"series:                {len(keys)}",
             f"series with 1 vintage: {single}",
             f"workbooks opened:      {books}",
             f"dates compared:        {dates:,}",
             f"values compared:       {values:,}",
             f"exactly equal:         {exact:,}",
             f"largest relative diff: {worst[0]:.3e}"]
    if worst[1]:
        lines.append(f"  at {worst[1]}")
    lines.append(f"mismatches over {TOLERANCE:g}: {len(bad)}")
    lines += [f"  {b}" for b in bad[:50]]

    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    LOG_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))
    print(f"log: {LOG_PATH.relative_to(ROOT).as_posix()}")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
